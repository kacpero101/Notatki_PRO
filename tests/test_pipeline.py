import copy
import plistlib
import subprocess

import pytest
from conftest import SAMPLE_GENERAL_NOTES, SAMPLE_NOTES, FakeProvider
from docx import Document

from lecture_notes import cli
from lecture_notes.config import load_config
from lecture_notes.modes import GENERAL, LECTURE
from lecture_notes.paths import notes_path_for, unique_path
from lecture_notes.pipeline import UserError, process_file, whisper_settings
from lecture_notes.quick_action import FILE_TYPES, install, shell_command, uninstall
from lecture_notes.transcript import Segment, Transcript


@pytest.fixture
def config(tmp_path, monkeypatch):
    monkeypatch.setenv("LECTURE_NOTES_CONFIG", str(tmp_path / "missing.toml"))
    return load_config()


def fake_transcriber(path, config):
    return Transcript([
        Segment("Dzień dobry, dziś pochodne.", 0.0, 3.0),
        Segment("To będzie na kolokwium.", 3.0, 6.0),
    ])


def test_unique_path(tmp_path):
    p = tmp_path / "a.txt"
    assert unique_path(p) == p
    p.write_text("x")
    assert unique_path(p).name == "a (2).txt"
    (tmp_path / "a (2).txt").write_text("x")
    assert unique_path(p).name == "a (3).txt"


def test_notes_path_name(tmp_path):
    assert notes_path_for(tmp_path / "Wykład 1.m4a").name == "Wykład 1 – notatki.docx"


def test_process_recording_writes_transcripts_and_docx(tmp_path, config):
    rec = tmp_path / "Wykład 1.m4a"
    rec.write_bytes(b"fake audio")
    provider = FakeProvider()
    result = process_file(rec, config, provider_factory=lambda cfg: provider, transcriber=fake_transcriber)

    txt, srt = result.transcript_paths
    assert txt.name == "Wykład 1.txt" and srt.name == "Wykład 1.srt"
    assert "To będzie na kolokwium." in txt.read_text()
    assert "00:00:03,000 --> 00:00:06,000" in srt.read_text()
    assert result.notes_path.name == "Wykład 1 – notatki.docx"
    assert result.cost == pytest.approx((1000 * 2 + 500 * 10) / 1e6)  # sonnet price from example config
    assert "[00:00:00]" in provider.calls[0]["user"]
    doc = Document(str(result.notes_path))
    assert doc.paragraphs[0].text == "Wykład 1.m4a"


def test_never_overwrites(tmp_path, config):
    rec = tmp_path / "w.mp3"
    rec.write_bytes(b"x")
    for name in ("w.txt", "w.srt", "w – notatki.docx"):
        (tmp_path / name).write_text("stare")
    result = process_file(rec, config, provider_factory=lambda c: FakeProvider(), transcriber=fake_transcriber)
    assert result.notes_path.name == "w – notatki (2).docx"
    assert [p.name for p in result.transcript_paths] == ["w (2).txt", "w (2).srt"]
    for name in ("w.txt", "w.srt", "w – notatki.docx"):
        assert (tmp_path / name).read_text() == "stare"


def test_txt_input_without_timestamps(tmp_path, config):
    src = tmp_path / "notatki.txt"
    src.write_text("Wykład bez znaczników czasu.\n\nDrugi akapit.", encoding="utf-8")
    notes = copy.deepcopy(SAMPLE_NOTES)
    provider = FakeProvider(responses=[notes])
    result = process_file(src, config, provider_factory=lambda c: provider,
                          transcriber=lambda *a: pytest.fail("must not transcribe text input"))
    assert result.transcript_paths is None
    text = "\n".join(p.text for p in Document(str(result.notes_path)).paragraphs)
    assert "[00:" not in text  # model's stamps were dropped


def test_srt_input(tmp_path, config):
    src = tmp_path / "w.srt"
    src.write_text("1\n00:00:01,000 --> 00:00:02,000\nTreść.\n", encoding="utf-8")
    result = process_file(src, config, provider_factory=lambda c: FakeProvider(), transcriber=None)
    assert result.notes_path.exists()


@pytest.mark.parametrize("name, content, needle", [
    ("a.pdf", b"x", "Nieobsługiwany"),
    ("pusty.txt", b"   \n", "Pusta"),
])
def test_user_errors(tmp_path, config, name, content, needle):
    f = tmp_path / name
    f.write_bytes(content)
    with pytest.raises(UserError, match=needle):
        process_file(f, config, provider_factory=lambda c: FakeProvider())


def test_missing_file(tmp_path, config):
    with pytest.raises(UserError, match="Nie znaleziono"):
        process_file(tmp_path / "brak.m4a", config, provider_factory=lambda c: FakeProvider())


def test_silent_recording(tmp_path, config):
    rec = tmp_path / "cisza.wav"
    rec.write_bytes(b"x")
    with pytest.raises(UserError, match="Nie rozpoznano"):
        process_file(rec, config, provider_factory=lambda c: FakeProvider(), transcriber=lambda *a: Transcript([]))


def test_cli_notifies_start_success_and_error(tmp_path, config, monkeypatch):
    good = tmp_path / "w.srt"
    good.write_text("1\n00:00:01,000 --> 00:00:02,000\nTreść.\n", encoding="utf-8")
    bad = tmp_path / "x.pdf"
    bad.write_bytes(b"x")
    sent = []
    monkeypatch.setattr(cli, "notify", lambda msg, title="", sound=None: sent.append((title, msg)))
    monkeypatch.setattr(cli, "LOG_DIR", tmp_path / "logs")
    monkeypatch.setattr("lecture_notes.pipeline.create_provider", lambda cfg: FakeProvider())
    code = cli.main([str(good), str(bad), "-q"])
    assert code == 1
    msgs = [m for _, m in sent]
    assert msgs[0].startswith("Przetwarzam (1/2): w.srt")
    assert any(m.startswith("Gotowe: w – notatki.docx") and "Koszt API ≈ $" in m for m in msgs)
    assert any("Nieobsługiwany" in m for t, m in sent if "błąd" in t)
    assert (tmp_path / "logs" / "lecture-notes.log").exists()


def _workflow(target):
    info = plistlib.loads((target / "Contents" / "Info.plist").read_bytes())
    wflow = plistlib.loads((target / "Contents" / "document.wflow").read_bytes())
    return info["NSServices"][0], wflow["actions"][0]["action"]["ActionParameters"]


def test_quick_action_bundles(tmp_path):
    exe = tmp_path / "bin" / "lecture-notes"
    targets = install(exe, services_dir=tmp_path)
    assert [t.name for t in targets] == ["Zrób notatki.workflow", "Zrób notatki (dowolny temat).workflow"]
    for target in targets:
        for name in ("Info.plist", "document.wflow"):
            lint = subprocess.run(["/usr/bin/plutil", "-lint", str(target / "Contents" / name)], capture_output=True)
            assert lint.returncode == 0, lint.stdout

    service, params = _workflow(targets[0])
    assert service["NSMenuItem"]["default"] == "Zrób notatki"
    assert service["NSSendFileTypes"] == FILE_TYPES
    assert params["inputMethod"] == 1 and params["shell"] == "/bin/zsh"
    assert str(exe) in params["COMMAND_STRING"] and "--mode" not in params["COMMAND_STRING"]
    assert 'export PATH="/opt/homebrew/bin' in params["COMMAND_STRING"]

    service, params = _workflow(targets[1])
    assert service["NSMenuItem"]["default"] == "Zrób notatki (dowolny temat)"
    assert f'{exe} --mode general "$@"' in params["COMMAND_STRING"]

    assert uninstall(services_dir=tmp_path) == 2
    assert not any(t.exists() for t in targets)
    assert uninstall(services_dir=tmp_path) == 0


def test_shell_command_quotes_paths_with_spaces(tmp_path):
    cmd = shell_command(tmp_path / "my dir" / "lecture-notes", ["--mode", "general"])
    assert "'" in cmd and '--mode general "$@"' in cmd


def test_general_mode_pipeline(tmp_path, config):
    rec = tmp_path / "Podcast.mp3"
    rec.write_bytes(b"x")
    seen = {}

    def transcriber(path, whisper_cfg):
        seen.update(whisper_cfg)
        return Transcript([Segment("Reinforcement learning is about rewards.", 0.0, 50.0)], language="en")

    provider = FakeProvider(responses=[copy.deepcopy(SAMPLE_GENERAL_NOTES)])
    result = process_file(rec, config, provider_factory=lambda c: provider, transcriber=transcriber, mode=GENERAL)

    assert seen["language"] == "" and seen["initial_prompt"] == ""  # auto-detect, no course hint
    assert seen["model"] == config["whisper"]["model"]
    assert provider.calls[0]["system"] == GENERAL.system_prompt
    assert "Wykryty język nagrania: en" in provider.calls[0]["user"]
    doc = Document(str(result.notes_path))
    headings = [p.text for p in doc.paragraphs if p.style.name == "Heading 1"]
    assert "Najważniejsze wnioski" in headings and "Na egzamin / zaliczenie" not in headings
    assert doc.core_properties.subject == "Notatki"


def test_lecture_mode_keeps_polish_whisper_settings(config):
    config["whisper"]["initial_prompt"] = "Analiza matematyczna"
    assert whisper_settings(config, LECTURE)["language"] == "pl"
    assert whisper_settings(config, LECTURE)["initial_prompt"] == "Analiza matematyczna"
    general = whisper_settings(config, GENERAL)
    assert general["language"] == "" and general["initial_prompt"] == ""
    config["modes"]["general"]["whisper_language"] = "en"
    assert whisper_settings(config, GENERAL)["language"] == "en"


def test_cli_mode_flag(tmp_path, config, monkeypatch):
    src = tmp_path / "film.srt"
    src.write_text("1\n00:00:01,000 --> 00:00:02,000\nContent.\n", encoding="utf-8")
    sent, used = [], []
    monkeypatch.setattr(cli, "notify", lambda msg, title="", sound=None: sent.append(title))
    monkeypatch.setattr(cli, "LOG_DIR", tmp_path / "logs")

    def fake_factory(cfg):
        provider = FakeProvider(responses=[copy.deepcopy(SAMPLE_GENERAL_NOTES)])
        used.append(provider)
        return provider

    monkeypatch.setattr("lecture_notes.pipeline.create_provider", fake_factory)
    assert cli.main(["--mode", "general", str(src), "-q"]) == 0
    assert used[0].calls[0]["system"] == GENERAL.system_prompt
    assert set(sent) == {"Notatki"}
