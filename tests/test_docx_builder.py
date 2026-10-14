from datetime import date, datetime

from conftest import SAMPLE_GENERAL_NOTES
from docx import Document

from lecture_notes.docx_builder import FORMULA_STYLE, MONO_FONT, DocumentMeta, save_document
from lecture_notes.modes import GENERAL
from lecture_notes.notes import empty_notes

META = DocumentMeta(
    title="Analiza 1 – wykład 3.m4a",
    recorded=date(2026, 10, 7),
    model="claude-sonnet-5-5",
    cost="≈ $0.12",
    generated=datetime(2026, 10, 11, 14, 5),
    source_name="Analiza 1 – wykład 3.m4a",
)

EXPECTED_H1 = [
    "Streszczenie", "Notatki tematyczne", "Definicje", "Wzory i przykłady",
    "Na egzamin / zaliczenie", "Pytania do powtórki", "Fragmenty do sprawdzenia", "Metadane",
]


def _build(tmp_path, notes):
    path = save_document(notes, META, tmp_path / "out.docx")
    return Document(str(path))


def _headings(doc, level):
    return [p.text for p in doc.paragraphs if p.style.name == f"Heading {level}"]


def test_sections_order_and_title(tmp_path, sample_notes):
    doc = _build(tmp_path, sample_notes)
    assert doc.paragraphs[0].style.name == "Title"
    assert doc.paragraphs[0].text == META.title
    assert "07.10.2026" in doc.paragraphs[1].text
    assert _headings(doc, 1) == EXPECTED_H1
    assert doc.core_properties.title == META.title


def test_content_rendering(tmp_path, sample_notes):
    doc = _build(tmp_path, sample_notes)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Pochodna  [00:00:03]" in _headings(doc, 2)
    assert "Prowadzący powiedział wprost" in _headings(doc, 2)
    assert "[00:00:03] Liczenie pochodnej wielomianu" in text
    assert "„To będzie na kolokwium.”" in text
    assert "Dlaczego ważne: Powtórzone kilka razy." in text
    assert "„pochodna z iksa” — Niejasne" in text

    formula = next(p for p in doc.paragraphs if p.style.name == FORMULA_STYLE)
    assert formula.text == "f(x) = x^2 + 1  =>  f'(x) = 2x"
    assert doc.styles[FORMULA_STYLE].font.name == MONO_FONT

    bullets = [p for p in doc.paragraphs if p.style.name == "List Bullet"]
    numbered = [p.text for p in doc.paragraphs if p.style.name == "List Number"]
    assert any("Definicja przez granicę" in p.text for p in bullets)
    assert numbered == ["Jak definiujemy pochodną?"]


def test_metadata_and_footer(tmp_path, sample_notes):
    doc = _build(tmp_path, sample_notes)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Model: claude-sonnet-5-5" in text
    assert "Koszt wywołania API: ≈ $0.12" in text
    assert "Wygenerowano: 11.10.2026 14:05" in text
    assert "claude-sonnet-5-5" in doc.sections[0].footer.paragraphs[0].text


def test_empty_notes_still_have_all_sections(tmp_path):
    doc = _build(tmp_path, empty_notes())
    assert _headings(doc, 1) == EXPECTED_H1
    assert sum(p.text == "Brak." for p in doc.paragraphs) >= 6


def test_polish_language_set(tmp_path, sample_notes):
    doc = _build(tmp_path, sample_notes)
    xml = doc.styles["Normal"].element.xml
    assert 'w:val="pl-PL"' in xml


GENERAL_H1 = [
    "Streszczenie", "Notatki tematyczne", "Kluczowe pojęcia", "Najważniejsze wnioski",
    "Pytania do powtórki", "Uzupełnienia (spoza nagrania)", "Fragmenty do sprawdzenia", "Metadane",
]


def test_general_mode_sections(tmp_path):
    path = save_document(SAMPLE_GENERAL_NOTES, META, tmp_path / "g.docx", GENERAL)
    doc = Document(str(path))
    # formulas are optional and empty here -> section omitted
    assert _headings(doc, 1) == GENERAL_H1
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "[00:00:45] Nagroda definiuje cel agenta." in text
    assert "Kontekst — Termin pochodzi z psychologii behawioralnej." in text
    assert "Tych informacji nie ma w nagraniu" in text
    assert "Na egzamin" not in text


def test_general_mode_optional_sections_shown_when_present(tmp_path):
    notes = dict(SAMPLE_GENERAL_NOTES, formulas=[
        {"description": "Zysk", "formula": "G = sum_{t} gamma^t r_t", "example": "", "timestamp": ""},
    ], supplements=[])
    doc = Document(str(save_document(notes, META, tmp_path / "g.docx", GENERAL)))
    headings = _headings(doc, 1)
    assert "Wzory i przykłady" in headings
    assert "Uzupełnienia (spoza nagrania)" not in headings
