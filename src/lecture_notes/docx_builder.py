"""Build the Word document from structured notes (python-docx, built-in styles)."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph

from .modes import LECTURE, NoteMode

MONO_FONT = "Menlo"
FORMULA_STYLE = "Wzór"
GRAY = RGBColor(0x6B, 0x6B, 0x6B)


@dataclass
class DocumentMeta:
    title: str
    recorded: date | None
    model: str
    cost: str
    generated: datetime
    source_name: str


# ------------------------------------------------------------------ helpers

def _set_language(style, lang: str = "pl-PL") -> None:
    rpr = style.element.get_or_add_rPr()
    el = rpr.find(qn("w:lang"))
    if el is None:
        el = OxmlElement("w:lang")
        rpr.append(el)
    el.set(qn("w:val"), lang)


def _prepare_styles(doc) -> None:
    styles = doc.styles
    _set_language(styles["Normal"])
    if FORMULA_STYLE not in [s.name for s in styles]:
        st = styles.add_style(FORMULA_STYLE, WD_STYLE_TYPE.PARAGRAPH)
        st.base_style = styles["Normal"]
        st.font.size = Pt(10.5)
        rfonts = st.element.get_or_add_rPr().get_or_add_rFonts()
        for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
            rfonts.set(qn(attr), MONO_FONT)
        st.paragraph_format.left_indent = Pt(18)
        st.paragraph_format.space_after = Pt(4)


def _timestamp(paragraph: Paragraph, ts: str) -> None:
    if ts:
        run = paragraph.add_run(f"[{ts}] ")
        run.font.color.rgb = GRAY
        run.font.size = Pt(9)


def _mono_run(paragraph: Paragraph, text: str):
    run = paragraph.add_run(text)
    run.font.name = MONO_FONT
    run.font.size = Pt(10.5)
    rfonts = run._element.get_or_add_rPr().get_or_add_rFonts()
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), MONO_FONT)
    return run


def _none(doc, text: str = "Brak.") -> None:
    doc.add_paragraph().add_run(text).italic = True


def _items(notes: dict[str, Any], key: str) -> list[Any]:
    value = notes.get(key) or []
    return value if isinstance(value, list) else []


# ----------------------------------------------------------------- sections
# Each renderer takes (doc, notes, heading); the order comes from the note mode.

def _summary(doc, notes, heading):
    doc.add_heading(heading, level=1)
    if notes.get("summary"):
        doc.add_paragraph(notes["summary"])
    else:
        _none(doc)


def _topics(doc, notes, heading):
    doc.add_heading(heading, level=1)
    topics = _items(notes, "topics")
    if not topics:
        _none(doc)
    for topic in topics:
        h = doc.add_heading(level=2)
        h.add_run(topic.get("title", ""))
        if topic.get("timestamp"):
            run = h.add_run(f"  [{topic['timestamp']}]")
            run.font.size = Pt(9)
            run.font.color.rgb = GRAY
        for point in topic.get("points", []):
            doc.add_paragraph(point, style="List Bullet")


def _definitions(doc, notes, heading):
    doc.add_heading(heading, level=1)
    defs = _items(notes, "definitions")
    if not defs:
        _none(doc)
    for d in defs:
        p = doc.add_paragraph(style="List Bullet")
        _timestamp(p, d.get("timestamp", ""))
        p.add_run(d.get("term", "")).bold = True
        p.add_run(" — " + d.get("definition", ""))


def _takeaways(doc, notes, heading):
    doc.add_heading(heading, level=1)
    items = _items(notes, "takeaways")
    if not items:
        _none(doc)
    for t in items:
        p = doc.add_paragraph(style="List Bullet")
        _timestamp(p, t.get("timestamp", ""))
        p.add_run(t.get("text", ""))


def _formulas(doc, notes, heading):
    doc.add_heading(heading, level=1)
    formulas = _items(notes, "formulas")
    if not formulas:
        _none(doc)
    for f in formulas:
        p = doc.add_paragraph()
        _timestamp(p, f.get("timestamp", ""))
        p.add_run(f.get("description", "")).bold = True
        doc.add_paragraph(f.get("formula", ""), style=FORMULA_STYLE)
        if f.get("example"):
            ex = doc.add_paragraph()
            ex.paragraph_format.left_indent = Pt(18)
            ex.add_run("Przykład: ").italic = True
            _mono_run(ex, f["example"])


def _exam(doc, notes, heading):
    doc.add_heading(heading, level=1)
    items = _items(notes, "exam_items")
    groups = [
        ("Prowadzący powiedział wprost", [i for i in items if i.get("kind") == "explicit"]),
        ("Prawdopodobne (akcentowane lub powtarzane)", [i for i in items if i.get("kind") != "explicit"]),
    ]
    for group_heading, group in groups:
        doc.add_heading(group_heading, level=2)
        if not group:
            _none(doc)
        for item in group:
            p = doc.add_paragraph(style="List Bullet")
            _timestamp(p, item.get("timestamp", ""))
            p.add_run(item.get("content", "")).bold = True
            if item.get("quote"):
                p.add_run("\nCytat: ").italic = True
                p.add_run(f"„{item['quote']}”").italic = True
            if item.get("reason"):
                p.add_run("\nDlaczego ważne: ").italic = True
                p.add_run(item["reason"])


def _questions(doc, notes, heading):
    doc.add_heading(heading, level=1)
    questions = _items(notes, "review_questions")
    if not questions:
        _none(doc)
    for q in questions:
        doc.add_paragraph(q, style="List Number")


def _supplements(doc, notes, heading):
    doc.add_heading(heading, level=1)
    note = doc.add_paragraph().add_run(
        "Tych informacji nie ma w nagraniu — dodał je model językowy jako wyjaśnienie lub kontekst. "
        "Mogą zawierać błędy; traktuj je jako wskazówkę, nie jako treść źródła."
    )
    note.italic = True
    note.font.size = Pt(9)
    note.font.color.rgb = GRAY
    items = _items(notes, "supplements")
    if not items:
        _none(doc)
    for s in items:
        p = doc.add_paragraph(style="List Bullet")
        p.add_run(s.get("title", "")).bold = True
        if s.get("content"):
            p.add_run(" — " + s["content"])


def _issues(doc, notes, heading):
    doc.add_heading(heading, level=1)
    issues = _items(notes, "transcript_issues")
    if not issues:
        _none(doc, "Nie wykryto podejrzanych fragmentów.")
    for i in issues:
        p = doc.add_paragraph(style="List Bullet")
        _timestamp(p, i.get("timestamp", ""))
        p.add_run(f"„{i.get('fragment', '')}”").bold = True
        if i.get("problem"):
            p.add_run(" — " + i["problem"])
        if i.get("suggestion"):
            p.add_run("\nPrawdopodobnie: ").italic = True
            p.add_run(i["suggestion"])


RENDERERS = {
    "summary": _summary,
    "topics": _topics,
    "definitions": _definitions,
    "takeaways": _takeaways,
    "formulas": _formulas,
    "exam_items": _exam,
    "review_questions": _questions,
    "supplements": _supplements,
    "transcript_issues": _issues,
}


def _metadata(doc, meta: DocumentMeta):
    doc.add_heading("Metadane", level=1)
    lines = [
        f"Źródło: {meta.source_name}",
        f"Model: {meta.model}",
        f"Koszt wywołania API: {meta.cost}",
        f"Wygenerowano: {meta.generated:%d.%m.%Y %H:%M}",
        "Notatki wygenerowane automatycznie z transkrypcji — sprawdź fragmenty oznaczone powyżej.",
    ]
    for line in lines:
        p = doc.add_paragraph()
        run = p.add_run(line)
        run.font.size = Pt(9)
        run.font.color.rgb = GRAY

    footer = doc.sections[0].footer.paragraphs[0]
    run = footer.add_run(
        f"lecture-notes · {meta.model} · koszt {meta.cost} · {meta.generated:%d.%m.%Y %H:%M}"
    )
    run.font.size = Pt(8)
    run.font.color.rgb = GRAY


# --------------------------------------------------------------------- main

def build_document(notes: dict[str, Any], meta: DocumentMeta, mode: NoteMode = LECTURE):
    doc = Document()
    _prepare_styles(doc)

    doc.core_properties.title = meta.title
    doc.core_properties.subject = mode.label
    doc.core_properties.author = "lecture-notes"
    doc.core_properties.language = "pl-PL"

    doc.add_heading(meta.title, level=0)
    sub = doc.add_paragraph()
    if meta.recorded:
        sub.add_run(f"Data nagrania: {meta.recorded:%d.%m.%Y}").italic = True
    else:
        sub.add_run("Data nagrania: nieznana").italic = True

    for section in mode.sections:
        if section.optional and not notes.get(section.key):
            continue
        RENDERERS[section.key](doc, notes, section.heading)
    _metadata(doc, meta)
    return doc


def save_document(notes: dict[str, Any], meta: DocumentMeta, path: Path, mode: NoteMode = LECTURE) -> Path:
    build_document(notes, meta, mode).save(str(path))
    return path
